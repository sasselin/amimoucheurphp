# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Transforme catalogue.json (sortie de export.php) en CSV et XLSX faciles à travailler.

Usage : uv run export-catalogue/build.py <catalogue.json> <dossier_sortie>
"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

XLSX_CELL_MAX = 32767
MAX_ATTRS = 3


def t(d, lang="fr"):
    """Texte d'un champ multilingue {fr, en}."""
    return (d or {}).get(lang) or ""


def join(items, sep=" | "):
    return sep.join(str(i) for i in items if i not in (None, ""))


def promo_cols(promo):
    promo = promo or {}
    return {
        "promo_type": promo.get("type", ""),
        "promo_valeur": promo.get("value", ""),
        "promo_rabais_pct": promo.get("discount_pct", ""),
        "promo_source": promo.get("source", ""),
        "promo_regle": promo.get("rule") or "",
        "promo_debut": promo.get("start") or "",
        "promo_fin": promo.get("end") or "",
    }


def cover(p):
    return p["images"][0]["url"] if p["images"] else ""  # la principale est en premier


def categories_cols(p):
    dc = p.get("default_category") or {}
    return {
        "marque": p["brand"] or "",
        "categorie_principale": t(dc.get("path")),
        "categories": join([t(c["path"]) for c in p["categories"]]),
        "categories_ids": join([c["id"] for c in p["categories"]], ","),
    }


def variant_rows(doc):
    """Une ligne par article vendable : chaque variante, ou le produit lui-même s'il n'en a pas."""
    rows = []
    for p in doc["products"]:
        for v in p["variants"] or [None]:
            price = (v or p)["price"]
            row = {
                "produit_id": p["id"],
                "variante_id": v["id"] if v else "",
                "sku": (v and v["reference"]) or p["reference"] or "",
                "nom_fr": t(p["name"]),
                "nom_en": t(p["name"], "en"),
                "variante_fr": t(v["label"]) if v else "",
                "variante_en": t(v["label"], "en") if v else "",
            }
            attrs = v["attributes"] if v else []
            for n in range(MAX_ATTRS):
                a = attrs[n] if n < len(attrs) else None
                row[f"attribut{n+1}"] = t(a["group"]) if a else ""
                row[f"valeur{n+1}"] = t(a["value"]) if a else ""
            row.update({
                "prix_regulier_ht": price["regular"],
                "prix_final_ht": price["final"],
                "en_promo": price["promo"] is not None,
            })
            row.update(promo_cols(price["promo"]))
            row.update({
                "quantite": v["quantity"] if v else p["stock"]["quantity"],
                "vente_si_rupture": p["stock"]["sell_when_out_of_stock"],
                "poids_kg": v["weight"] if v else p["weight"],
            })
            row.update(categories_cols(p))
            row.update({
                "image": (v and v["images"] and v["images"][0]) or cover(p),
                "url_fr": p["seo"]["fr"]["url"],
            })
            rows.append(row)
    return rows


def product_rows(doc):
    rows = []
    for p in doc["products"]:
        pr = p["price"]
        row = {
            "produit_id": p["id"],
            "reference": p["reference"] or "",
            "nom_fr": t(p["name"]),
            "nom_en": t(p["name"], "en"),
        }
        row.update(categories_cols(p))
        row.update({
            "nb_variantes": len(p["variants"]),
            "prix_regulier_ht": pr["regular"],
            "prix_final_ht": pr["final"],
            "prix_min_ht": pr.get("min", pr["final"]),
            "prix_max_ht": pr.get("max", pr["final"]),
            "en_promo": pr["promo"] is not None or pr.get("some_variants_on_promo", False),
            "quantite": p["stock"]["quantity"],
            "vente_si_rupture": p["stock"]["sell_when_out_of_stock"],
            "derniere_vente": p["last_sale"] or "",
            "nb_images": len(p["images"]),
            "image_principale": cover(p),
            "tags_fr": join((p["tags"] or {}).get("fr", []), ", "),
            "tags_en": join((p["tags"] or {}).get("en", []), ", "),
            "url_fr": p["seo"]["fr"]["url"],
            "url_en": p["seo"]["en"]["url"],
            "meta_titre_fr": p["seo"]["fr"]["meta_title"] or "",
            "meta_titre_en": p["seo"]["en"]["meta_title"] or "",
            "meta_description_fr": p["seo"]["fr"]["meta_description"] or "",
            "meta_description_en": p["seo"]["en"]["meta_description"] or "",
            "description_courte_fr": t(p["description_short"]),
            "description_courte_en": t(p["description_short"], "en"),
            "description_fr": t(p["description"]),
            "description_en": t(p["description"], "en"),
            "date_creation": p["date_add"] or "",
            "date_modif": p["date_upd"] or "",
        })
        if "excluded_reasons" in p:
            row["exclu_car"] = join(p["excluded_reasons"], ", ")
        rows.append(row)
    return rows


def category_rows(doc):
    rows = []
    for c in doc["categories"]:
        row = {
            "categorie_id": c["id"],
            "parent_id": c["parent_id"] or "",
            "profondeur": c["depth"],
            "position": c["position"],
            "nom_fr": t(c["name"]),
            "nom_en": t(c["name"], "en"),
            "chemin_fr": t(c["path"]),
            "chemin_en": t(c["path"], "en"),
            "nb_produits": c["product_count"],
            "nb_produits_avec_sous_categories": c["product_count_with_children"],
            "url_fr": c["seo"]["fr"]["url"],
            "url_en": c["seo"]["en"]["url"],
            "meta_titre_fr": c["seo"]["fr"]["meta_title"] or "",
            "meta_description_fr": c["seo"]["fr"]["meta_description"] or "",
        }
        if "excluded_reasons" in c:
            row["exclue_car"] = join(c["excluded_reasons"], ", ")
        rows.append(row)
    return rows


def image_rows(doc):
    rows = []
    for p in doc["products"]:
        for n, img in enumerate(p["images"], 1):
            rows.append({
                "produit_id": p["id"],
                "nom_fr": t(p["name"]),
                "ordre": n,
                "principale": img["cover"],
                "url_originale": img["url"],
                "legende_fr": t(img["legend"]),
                "variantes_ids": join([v["id"] for v in p["variants"] if img["url"] in v["images"]], ","),
                "fichier_manquant": img["file_missing"],
            })
    return rows


def promo_rows(doc):
    """Promos en cours seulement, une ligne par article touché."""
    return [
        {k: r[k] for k in ("produit_id", "variante_id", "sku", "nom_fr", "variante_fr", "prix_regulier_ht",
                           "prix_final_ht", "promo_type", "promo_valeur", "promo_rabais_pct", "promo_source",
                           "promo_regle", "promo_debut", "promo_fin", "quantite")}
        for r in variant_rows(doc) if r["en_promo"]
    ]


def url_rows(doc):
    rows = []
    for kind, items in (("produit", doc["products"]), ("categorie", doc["categories"])):
        for x in items:
            for lang, s in x["seo"].items():
                rows.append({"type": kind, "id": x["id"], "langue": lang, "nom": t(x["name"], lang),
                             "ancienne_url": s["url"], "nouvelle_url": ""})
    return rows


def write_csv(path, rows):
    if not rows:
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def add_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    if not rows:
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    truncated = 0
    for r in rows:
        values = []
        for h in headers:
            v = r[h]
            if isinstance(v, str) and len(v) > XLSX_CELL_MAX:
                v = v[: XLSX_CELL_MAX - 20] + " …[TRONQUÉ]"
                truncated += 1
            values.append(v)
        ws.append(values)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        width = min(60, max(10, len(h) + 2, *(len(str(r[h])) for r in rows[:200])))
        ws.column_dimensions[get_column_letter(i)].width = width
    if truncated:
        print(f"  {title} : {truncated} cellule(s) tronquée(s) à {XLSX_CELL_MAX} caractères (texte complet dans le JSON/CSV)")


def write_html(path, doc):
    """Mini catalogue autonome : gabarit view.html + données intégrées."""
    template = (Path(__file__).parent / "view.html").read_text(encoding="utf-8")
    data = json.dumps(doc, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    path.write_text(template.replace("/*__DATA__*/", data), encoding="utf-8")


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    doc = json.loads(src.read_text(encoding="utf-8"))
    if src.resolve() != (out / "catalogue.json").resolve():
        (out / "catalogue.json").write_text(src.read_text(encoding="utf-8"), encoding="utf-8")

    sheets = {
        "Produits": product_rows(doc),
        "Variantes": variant_rows(doc),
        "Categories": category_rows(doc),
        "Images": image_rows(doc),
        "Promotions": promo_rows(doc),
        "URL": url_rows(doc),
    }
    write_csv(out / "produits.csv", sheets["Produits"])
    write_csv(out / "variantes.csv", sheets["Variantes"])
    write_csv(out / "categories.csv", sheets["Categories"])
    write_csv(out / "url_mapping.csv", sheets["URL"])

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        add_sheet(wb, title, rows)
    wb.save(out / "catalogue.xlsx")
    write_html(out / "catalogue.html", doc)

    m = doc["meta"]
    c = m["counts"]
    print(f"Sortie : {out}")
    print(f"  {c['products']} produits, {c['variants']} variantes, {c['categories']} catégories")
    print("  Exclus : " + ", ".join(f"{k}={v}" for k, v in m["excluded"].items()))
    print("  " + ", ".join(f"{k}: {len(v)} lignes" for k, v in sheets.items()))


if __name__ == "__main__":
    main()
